"""템플릿 docx에서 실제 근무 데이터(실명·결함내용·문서 메타데이터)를 지운다.

업무일지 원본을 서식 템플릿으로 쓰기 때문에, 저장된 파일 자체에 지난 근무의
인수인계자 실명·IRR 결함 내용·작업 사항이 그대로 남는다. 저장소에 올리기 전에
반드시 한 번 돌려서 빈 서식으로 만든다.
"""

from __future__ import annotations

import sys

import docx

from worksheet_auto.docx_writer import (
    OUTSTATION_HEADING,
    SECTION_HEADINGS,
    _section_content_row,
    clear_carryover,
    set_cell_text,
    set_paragraph_text,
)


def scrub(path: str, out_path: str | None = None) -> str:
    out_path = out_path or path
    document = docx.Document(path)
    table = document.tables[0]

    clear_carryover(table)
    for heading in list(SECTION_HEADINGS.values()) + [OUTSTATION_HEADING]:
        row = _section_content_row(table, heading)
        if row is not None:
            set_cell_text(table.rows[row].cells[0], "")

    for paragraph in document.paragraphs:
        if paragraph.text.strip().upper().startswith("DATE"):
            set_paragraph_text(paragraph, "DATE :")
            break

    props = document.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = ""
    props.subject = ""
    props.comments = ""
    props.category = ""
    props.keywords = ""
    props.revision = 1

    document.save(out_path)
    return out_path


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else "template/정비통제업무일지_TEMPLATE.docx"
    print("scrubbed:", scrub(target))

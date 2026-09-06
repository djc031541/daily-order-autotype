"""정비통제업무일지 docx 생성.

기존 업무일지 파일을 서식 템플릿으로 그대로 쓴다.
- 날짜가 바뀌면 틀리게 되는 칸(인수인계자/기상/FLT PLAN/근무인원/IRR/SEAT BLOCK/국내.해외)은 비우고
- '항공기 주요 작업 사항 (정비 1) / (정비 2)' 칸에 Daily Work Order 초안을 채운다.
서식(대명체 8pt Bold 등)은 템플릿의 기존 run 속성을 복제해 유지한다.
"""

from __future__ import annotations

import copy
import datetime as _dt
from typing import Iterable

import docx
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from .formatter import build_sections, render_section

WEEKDAY_KO = ["월", "화", "수", "목", "금", "토", "일"]

SECTION_HEADINGS = {
    "정비 1": "항공기 주요 작업 사항 (정비 1)",
    "정비 2": "항공기 주요 작업 사항 (정비 2)",
}
OUTSTATION_HEADING = "항공기 주요 작업 사항 (국내. 해외)"
IRR_HEADING_PREFIX = "항공기 IRR"
SKD_HEADING = "SKD MAINT"
SEAT_HEADING = "SEAT BLOCK 현황"


# --------------------------------------------------------------------------- #
# 서식을 유지한 채 텍스트를 갈아끼우는 저수준 헬퍼
# --------------------------------------------------------------------------- #
def _proto_rpr(container):
    for para in container.paragraphs:
        for run in para.runs:
            if run._element.rPr is not None:
                return copy.deepcopy(run._element.rPr)
    return None


def _clear_runs(paragraph: Paragraph) -> None:
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)


def _add_run(paragraph: Paragraph, text: str, rpr) -> None:
    run = paragraph.add_run(text)
    if rpr is not None:
        run._element.insert(0, copy.deepcopy(rpr))


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    rpr = None
    for run in paragraph.runs:
        if run._element.rPr is not None:
            rpr = copy.deepcopy(run._element.rPr)
            break
    _clear_runs(paragraph)
    _add_run(paragraph, text, rpr)


def set_cell_text(cell: _Cell, text: str) -> None:
    """셀 내용을 text 로 교체. 줄바꿈은 문단 분리, 기존 서식 유지."""
    rpr = _proto_rpr(cell)
    paragraphs = cell.paragraphs
    first = paragraphs[0]
    for para in paragraphs[1:]:
        para._p.getparent().remove(para._p)
    _clear_runs(first)

    lines = text.split("\n") if text else [""]
    _add_run(first, lines[0], rpr)

    prev = first._p
    for line in lines[1:]:
        new_p = copy.deepcopy(first._p)
        for run_el in new_p.findall(qn("w:r")):
            new_p.remove(run_el)
        prev.addnext(new_p)
        prev = new_p
        _add_run(Paragraph(new_p, first._parent), line, rpr)


def _unique_cells(row) -> "list[tuple[int, _Cell]]":
    """가로 병합된 셀 중복을 제거한 (grid index, cell) 목록."""
    seen = set()
    out = []
    for idx, cell in enumerate(row.cells):
        if cell._tc in seen:
            continue
        seen.add(cell._tc)
        out.append((idx, cell))
    return out


def _cell_at(row, index: int) -> _Cell | None:
    for idx, cell in _unique_cells(row):
        if idx == index:
            return cell
    return None


def _row_label(row) -> str:
    return " ".join(row.cells[0].text.split()).strip()


# --------------------------------------------------------------------------- #
# 템플릿 구조 탐색
# --------------------------------------------------------------------------- #
def _find_row(table: Table, predicate) -> int | None:
    for i, row in enumerate(table.rows):
        if predicate(_row_label(row)):
            return i
    return None


def _section_content_row(table: Table, heading: str) -> int | None:
    idx = _find_row(table, lambda t: t == heading)
    return idx + 1 if idx is not None and idx + 1 < len(table.rows) else None


# --------------------------------------------------------------------------- #
# 지난 일지 내용 비우기
# --------------------------------------------------------------------------- #
def _blank_cells(row, indices: Iterable[int]) -> None:
    for index in indices:
        cell = _cell_at(row, index)
        if cell is not None:
            set_cell_text(cell, "")


def _blank_row(row) -> None:
    for _, cell in _unique_cells(row):
        set_cell_text(cell, "")


def clear_carryover(table: Table) -> None:
    """전날 일지에서 남아 있으면 안 되는 값들을 비운다."""
    rows = table.rows

    for i, row in enumerate(rows):
        label = _row_label(row)

        # 인계자 / 인수자 / 기상
        if label == "인계자":
            _blank_cells(row, [2, 7, 14])
        # FLT PLAN 편수
        elif label.startswith("FLT PLAN"):
            _blank_cells(row, [0, 9, 15])
            first = _cell_at(row, 0)
            if first is not None:
                set_cell_text(first, "FLT PLAN")
        # 주간 / 야간 근무 인원
        elif label.replace(" ", "") in {"주간근무인원", "야간근무인원"}:
            _blank_cells(row, [6, 15])

    # IRR 데이터 행 (헤더 2줄 아래 ~ SKD MAINT 직전)
    irr = _find_row(table, lambda t: t.startswith(IRR_HEADING_PREFIX))
    skd = _find_row(table, lambda t: t == SKD_HEADING)
    if irr is not None and skd is not None:
        for i in range(irr + 2, skd):
            _blank_row(rows[i])
        # SKD MAINT 내용 행
        seat = _find_row(table, lambda t: t.endswith(SEAT_HEADING))
        stop = seat if seat is not None else skd + 2
        for i in range(skd + 1, stop):
            _blank_row(rows[i])

    # SEAT BLOCK 데이터 행 (헤더 2줄 아래 ~ 다음 섹션 제목 직전)
    seat = _find_row(table, lambda t: t.endswith(SEAT_HEADING))
    nxt = _find_row(table, lambda t: t.startswith("항공기 주요 작업 사항"))
    if seat is not None and nxt is not None:
        for i in range(seat + 2, nxt):
            _blank_row(rows[i])

    # 국내/해외 (IRR 성격 - 초안 대상 아님)
    out_row = _section_content_row(table, OUTSTATION_HEADING)
    if out_row is not None:
        _blank_row(rows[out_row])


# --------------------------------------------------------------------------- #
# 공개 API
# --------------------------------------------------------------------------- #
def format_date_line(date: _dt.date | None, shift: str) -> str:
    if date is None:
        return f"DATE : {shift}".strip()
    return f"DATE : {date:%Y}. {date:%m}. {date:%d}. ({WEEKDAY_KO[date.weekday()]}) {shift}".strip()


def output_filename(date: _dt.date | None, shift: str) -> str:
    stamp = date.strftime("%d-%b-%y").upper() if date else "UNKNOWN"
    return f"정비통제업무일지_({stamp}){shift}.docx"


def build_worksheet(
    template_path: str,
    out_path: str,
    entries,
    date: _dt.date | None,
    shift: str,
    arrow: bool = True,
    clear_previous: bool = True,
) -> "dict[str, str]":
    """템플릿을 열어 날짜와 정비1/정비2 초안을 채우고 out_path 로 저장.

    반환값: {섹션명: 채워넣은 본문} (미리보기/검증용)
    """
    document = docx.Document(template_path)
    table = document.tables[0]

    if clear_previous:
        clear_carryover(table)

    # 제목 아래 DATE 줄
    for paragraph in document.paragraphs:
        if paragraph.text.strip().upper().startswith("DATE"):
            set_paragraph_text(paragraph, format_date_line(date, shift))
            break

    sections = build_sections(entries)
    rendered: dict[str, str] = {}
    for name, heading in SECTION_HEADINGS.items():
        text = render_section(sections.get(name, {}), arrow=arrow)
        rendered[name] = text
        row_idx = _section_content_row(table, heading)
        if row_idx is not None:
            set_cell_text(table.rows[row_idx].cells[0], text)

    document.save(out_path)
    return rendered
